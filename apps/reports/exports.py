import csv
import re
from decimal import Decimal
from html import escape
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def _format_value(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _filename(report_type, file_format):
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M%S")
    return f"sportswear-{report_type}-{stamp}.{file_format}"


def export_csv(report_type, title, columns, rows):
    buffer = StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow([title])
    writer.writerow([f"Generado: {timezone.localtime():%Y-%m-%d %H:%M %Z}"])
    writer.writerow([label for _key, label in columns])
    for row in rows:
        writer.writerow([_format_value(row.get(key)) for key, _label in columns])
    response = HttpResponse("\ufeff" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report_type, "csv")}"'
    return response


def export_xlsx(report_type, title, columns, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "-", title)[:31] or "Reporte"
    sheet.append([title])
    sheet.append([f"Generado: {timezone.localtime():%Y-%m-%d %H:%M %Z}"])
    sheet.append([label for _key, label in columns])
    for row in rows:
        sheet.append([_format_value(row.get(key)) for key, _label in columns])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(columns)))
    sheet["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="18263A")
    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6759E8")
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(max(1, len(columns)))}{max(3, sheet.max_row)}"
    for column_index, (_key, label) in enumerate(columns, 1):
        values = [len(str(label))]
        values.extend(len(str(sheet.cell(row, column_index).value or "")) for row in range(4, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max(values, default=12) + 2, 12), 38)
    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(report_type, "xlsx")}"'
    return response


def export_pdf(report_type, title, columns, rows):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter), rightMargin=12 * mm,
        leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    content = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(escape(f"Generado: {timezone.localtime():%Y-%m-%d %H:%M %Z}"), styles["Normal"]),
        Spacer(1, 8 * mm),
    ]
    table_data = [[Paragraph(escape(label), styles["BodyText"]) for _key, label in columns]]
    for row in rows:
        table_data.append([
            Paragraph(escape(_format_value(row.get(key)) or "—"), styles["BodyText"])
            for key, _label in columns
        ])
    if len(table_data) == 1:
        table_data.append([Paragraph("Sin datos para los filtros seleccionados.", styles["BodyText"])] + [""] * (len(columns) - 1))
    table = Table(table_data, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18263A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDE3EA")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F9")]),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    content.append(table)
    doc.build(content)
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report_type, "pdf")}"'
    return response
